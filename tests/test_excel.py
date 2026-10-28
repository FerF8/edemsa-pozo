"""Las fórmulas del Excel tienen que dar el mismo reparto que el cálculo de referencia."""

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from pozo.excel_out import Sesion, escribir_reparto
from pozo.reparto import Lectura, repartir
from tests.formula_eval import Libro
from tests.test_reparto import _factura

D = Decimal


def _cerca(valor, esperado: Decimal) -> None:
    assert abs(D(str(valor)) - esperado) < D("0.001")


def test_las_formulas_dejan_el_consumo_faltante_sin_asignar(tmp_path: Path):
    factura = _factura()
    factura.kwh_alta = D("10")
    factura.kwh_baja = D("10")
    destino = tmp_path / "parcial.xlsx"
    escribir_reparto(
        factura,
        destino,
        personas=("A", "B", "C"),
        sesiones=[Sesion("A", D("0"), D("2"), D("0"), D("0"))],
    )
    referencia = repartir(
        factura,
        [Lectura("A", D("2"), D("0")), Lectura("B", D("0"), D("0")), Lectura("C", D("0"), D("0"))],
    )
    libro = Libro(load_workbook(destino))
    _cerca(libro.celda("Reparto", "B17"), referencia.personas[0].total)
    _cerca(libro.celda("Reparto", "C17"), referencia.personas[1].total)
    _cerca(libro.celda("Reparto", "D17"), referencia.personas[2].total)
    _cerca(libro.celda("Reparto", "E17"), referencia.sin_asignar.total)
    _cerca(libro.celda("Reparto", "B19"), D("0"))


def test_las_formulas_cierran_como_el_calculo(tmp_path: Path):
    factura = _factura()
    destino = tmp_path / "reparto.xlsx"
    escribir_reparto(
        factura,
        destino,
        personas=("A", "B", "C"),
        sesiones=[
            Sesion("A", D("0"), D("2"), D("0"), D("0")),
            Sesion("B", D("0"), D("1"), D("0"), D("2")),
        ],
    )
    referencia = repartir(
        factura,
        [Lectura("A", D("2"), D("0")), Lectura("B", D("1"), D("2")), Lectura("C", D("0"), D("0"))],
    )

    libro = Libro(load_workbook(destino))
    _cerca(libro.celda("Reparto", "B17"), referencia.personas[0].total)
    _cerca(libro.celda("Reparto", "C17"), referencia.personas[1].total)
    _cerca(libro.celda("Reparto", "D17"), referencia.personas[2].total)
    _cerca(libro.celda("Reparto", "E17"), referencia.sin_asignar.total)
    _cerca(libro.celda("Reparto", "B19"), D("0"))
